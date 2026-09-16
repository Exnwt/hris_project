import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.http import HttpResponse
import datetime

class DynamicExcelService:
    @staticmethod
    def import_data(file_obj, model_class, manual_mapping):
        wb = openpyxl.load_workbook(file_obj)
        ws = wb.active
        headers = [cell.value for cell in ws[1]] 
        
        # Cek index Excel mana yang cocok dengan manual mapping
        header_to_field = {}
        for idx, h in enumerate(headers):
            if h and h in manual_mapping:
                field_name = manual_mapping[h]
                try:
                    # Ambil object field asli dari model untuk validasi
                    field_obj = model_class._meta.get_field(field_name)
                    header_to_field[idx] = field_obj
                except Exception:
                    pass 

        success_count = 0
        errors = []

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row): continue

            row_dict = {}
            row_errors = []

            for idx, cell_val in enumerate(row):
                if idx not in header_to_field: continue
                
                field_obj = header_to_field[idx]
                field_name = field_obj.name

                if isinstance(cell_val, str) and cell_val.strip() == "":
                    cell_val = None

                # RELASI LOGIC
                if field_obj.is_relation and cell_val is not None:
                    related_model = field_obj.related_model
                    obj = None
                    lookup_fields = ['name', 'nama', 'nama_department', 'nama_company', 'nama_jabatan']
                    
                    for lookup in lookup_fields:
                        try:
                            obj = related_model.objects.get(**{f"{lookup}__iexact": str(cell_val).strip()})
                            break
                        except Exception:
                            continue
                    
                    if obj:
                        row_dict[field_name] = obj
                    else:
                        row_errors.append(f"Relasi '{cell_val}' untuk kolom '{field_obj.verbose_name.title()}' tidak ada di database.")
                else:
                    row_dict[field_name] = cell_val

            # VALIDASI REQUIRED FIELD (Khusus untuk field yang ada di manual mapping saja)
            for header_label, field_name in manual_mapping.items():
                f = model_class._meta.get_field(field_name)
                is_required = not f.blank and not f.null
                if is_required and (field_name not in row_dict or row_dict.get(field_name) is None):
                    row_errors.append(f"Kolom '{header_label}' wajib diisi.")

            if row_errors:
                errors.append(f"Baris {row_idx} Gagal: " + " | ".join(row_errors))
                continue

            try:
                model_class.objects.create(**row_dict)
                success_count += 1
            except Exception as e:
                errors.append(f"Baris {row_idx} Error DB: {str(e)}")

        return success_count, errors

    # -----------------------------------------------------
    # UPDATE FUNGSI TEMPLATE (MENGGUNAKAN MANUAL MAPPING)
    # -----------------------------------------------------
    @staticmethod
    def generate_template(model_class, manual_mapping, filename="Template_Import"):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Template Import"

        req_fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid") 
        opt_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid") 
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        inst_font = Font(name="Arial", size=9, italic=True, color="4B5563")
        inst_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

        headers = []
        instructions = []

        # LOOP HANYA BERDASARKAN MANUAL MAPPING YANG DITENTUKAN
        for header_label, field_name in manual_mapping.items():
            f = model_class._meta.get_field(field_name)
            is_required = not f.blank and not f.null
            
            inst_text = "Wajib diisi." if is_required else "Opsional."
            if f.is_relation:
                inst_text += f" (Isi NAMA {header_label}, BUKAN ID)"
                
            headers.append({"label": header_label, "required": is_required})
            instructions.append(inst_text)

        ws.append([h["label"] for h in headers])
        ws.append(instructions)

        for col_num, cell in enumerate(ws[1], 1):
            is_req = headers[col_num-1]["required"]
            cell.fill = req_fill if is_req else opt_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
            ws.cell(row=2, column=col_num).font = inst_font
            ws.cell(row=2, column=col_num).fill = inst_fill
            ws.cell(row=2, column=col_num).alignment = Alignment(wrap_text=True, vertical="top")

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
        wb.save(response)
        return response
    
    @staticmethod
    def export_data(queryset, selected_fields, filename="Export_Data"):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data Export"

        # Styling
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        border_style = Side(border_style="thin", color="CCCCCC")

        headers = [item['label'] for item in selected_fields]
        ws.append(headers)

        for col_num, cell in enumerate(ws[1], 1):
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(left=border_style, right=border_style, top=border_style, bottom=border_style)

        for obj in queryset:
            row_data = []
            for item in selected_fields:
                field_name = item['field']
                val = getattr(obj, field_name, None)

                # JIKA DATA ADALAH RELASI (ForeignKey) -> Ekstrak Namanya, Bukan ID
                if isinstance(val, models.Model):
                    if hasattr(val, 'name'): val = val.name
                    elif hasattr(val, 'nama'): val = val.nama
                    elif hasattr(val, 'nama_department'): val = val.nama_department
                    elif hasattr(val, 'nama_company'): val = val.nama_company
                    elif hasattr(val, 'nama_jabatan'): val = val.nama_jabatan
                    else: val = str(val) # Fallback ke method __str__ model tersebut

                if isinstance(val, (datetime.date, datetime.datetime)):
                    val = val.strftime("%Y-%m-%d")

                row_data.append(str(val) if val is not None else "")
            
            ws.append(row_data)

        # Auto Width
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
        wb.save(response)
        return response
    
    